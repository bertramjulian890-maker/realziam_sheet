from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from rwork_sheets_api import RworkSheetsClient


REQUIRED_HEADERS = ("店铺号", "店铺名称", "销售日期", "销售总金额")
DEFAULT_API_BASE_URL = "http://10.90.10.66:7055/"
DEFAULT_SPREADSHEET_TOKEN = "shtk99v6k5T9IvGMQjKnUbxZqub"


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    text = re.split(r"[ T]", text, maxsplit=1)[0].replace("年", "-").replace("月", "-").replace("日", "")
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def as_amount(value: Any) -> int | float:
    if value is None or str(value).strip() == "":
        raise ValueError("销售总金额为空")
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"销售总金额不是数字：{value}") from exc
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def filter_export(path: str | Path, business_date: date) -> list[dict[str, Any]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("导出表为空")
        positions = {str(value or "").strip(): index for index, value in enumerate(header)}
        missing = [name for name in REQUIRED_HEADERS if name not in positions]
        if missing:
            raise ValueError(f"导出表缺少字段：{', '.join(missing)}")

        result: list[dict[str, Any]] = []
        seen: set[str] = set()
        for row_number, row in enumerate(rows, start=2):
            sale_date = as_date(_cell(row, positions["销售日期"]))
            if sale_date != business_date:
                continue
            shop_id = str(_cell(row, positions["店铺号"]) or "").strip()
            if not shop_id:
                raise ValueError(f"第 {row_number} 行店铺号为空")
            if shop_id in seen:
                raise ValueError(f"重复店铺号：{shop_id}")
            seen.add(shop_id)
            result.append(
                {
                    "店铺号": shop_id,
                    "店铺名称": str(_cell(row, positions["店铺名称"]) or "").strip(),
                    "销售日期": business_date.isoformat(),
                    "销售总金额": as_amount(_cell(row, positions["销售总金额"])),
                }
            )
    finally:
        workbook.close()
    if not result:
        raise ValueError(f"导出表中没有 {business_date.isoformat()} 的销售记录")
    return result


def _cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def save_filtered_workbook(rows: list[dict[str, Any]], output: str | Path) -> Path:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "昨日销售"
    sheet.append(list(REQUIRED_HEADERS))
    for item in rows:
        sheet.append([item[name] for name in REQUIRED_HEADERS])
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 15
    sheet.column_dimensions["D"].width = 16
    workbook.save(output_path)
    return output_path


def column_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def header_date(value: Any, year: int) -> date | None:
    """Cloud date cells may be displayed text or spreadsheet serial numbers."""
    parsed = as_date(value)
    if parsed:
        return parsed
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=value)).date()
        except (OverflowError, ValueError):
            return None
    match = re.fullmatch(r"\s*(\d{1,2})\s*[月/-]\s*(\d{1,2})\s*日?\s*", str(value or ""))
    if match:
        try:
            return date(year, int(match[1]), int(match[2]))
        except ValueError:
            return None
    return None


def plan_updates(
    sheet_id: str,
    headers: list[list[Any]],
    cloud_ids: list[Any],
    rows: list[dict[str, Any]],
    business_date: date,
) -> list[dict[str, Any]]:
    clean_headers = [str(value or "").strip() for value in (headers[0] if headers else [])]
    if clean_headers[:2] != ["店铺号", "店铺名称"]:
        raise ValueError("云表格 A1、B1 必须分别为“店铺号”“店铺名称”")

    dates = headers[1] if len(headers) > 1 else []
    date_columns = []
    for index, value in enumerate(dates):
        parsed = header_date(value, business_date.year)
        if index >= 11 and parsed and (parsed.month, parsed.day) == (business_date.month, business_date.day):
            date_columns.append(index + 1)
    if len(date_columns) > 1:
        raise ValueError(f"云表格存在多个 {business_date.isoformat()} 日期列")
    updates: list[dict[str, Any]] = []
    if not date_columns:
        raise ValueError(f"第 2 行 L 列起找不到 {business_date.isoformat()} 日期列；固定模板不新增列")
    target_column = date_columns[0]

    id_to_row: dict[str, int] = {}
    for offset, value in enumerate(cloud_ids, start=3):
        shop_id = str(value or "").strip()
        if not shop_id or shop_id in {"合计", "总计", "小计", "求和"}:
            continue
        if shop_id in id_to_row:
            raise ValueError(f"云表格店铺号重复：{shop_id}")
        id_to_row[shop_id] = offset
    missing = [item["店铺号"] for item in rows if item["店铺号"] not in id_to_row]
    if missing:
        raise ValueError(f"以下店铺号在云表格中找不到：{', '.join(missing[:20])}")

    letter = column_letter(target_column)
    for item in rows:
        row_number = id_to_row[item["店铺号"]]
        updates.append({"range": f"{sheet_id}!{letter}{row_number}:{letter}{row_number}", "values": [[item["销售总金额"]]]})
    return updates


def plan_total(sheet_id: str, column: str, cloud_ids: list[Any]) -> dict[str, Any]:
    shop_rows = [index for index, value in enumerate(cloud_ids, start=3)
                 if str(value or "").strip()
                 and str(value).strip() not in {"合计", "总计", "小计", "求和"}]
    if not shop_rows:
        raise ValueError("没有可用于汇总的店铺行")
    last_row = max(shop_rows)
    total_cell = f"{column}{last_row + 1}"
    return {"range": f"{sheet_id}!{total_cell}:{total_cell}",
            "values": [[f"=SUM({column}3:{column}{last_row})"]]}


def sync_file(
    source_file: str | Path,
    business_date: date,
    *,
    api_base_url: str,
    spreadsheet_token: str,
    sheet_id: str = "",
    output_file: str | Path | None = None,
    max_rows: int = 2000,
    dry_run: bool = False,
) -> dict[str, Any]:
    rows = filter_export(source_file, business_date)
    filtered_path = save_filtered_workbook(
        rows,
        output_file or Path(source_file).with_name(f"{business_date.isoformat()}_昨日销售_筛选.xlsx"),
    )
    client = RworkSheetsClient(api_base_url)
    actual_sheet_id = sheet_id or client.sheet_id(spreadsheet_token)
    header_values = client.read_range(spreadsheet_token, f"{actual_sheet_id}!A1:ZZ2")
    id_values = client.read_range(spreadsheet_token, f"{actual_sheet_id}!A3:A{max_rows}")
    headers = header_values
    cloud_ids = [row[0] if row else "" for row in id_values]
    updates = plan_updates(actual_sheet_id, headers, cloud_ids, rows, business_date)
    target_column = re.search(r"!([A-Z]+)", updates[0]["range"])[1]
    total_update = plan_total(actual_sheet_id, target_column, cloud_ids)
    if not dry_run:
        for index in range(0, len(updates), 100):
            client.write_values(spreadsheet_token, updates[index : index + 100])
        written = client.read_range(spreadsheet_token, f"{actual_sheet_id}!{target_column}3:{target_column}{max_rows}")
        for update in updates:
            row_number = int(re.search(r"![A-Z]+(\d+)", update["range"])[1])
            actual = written[row_number - 3] if row_number - 3 < len(written) else []
            if not actual or as_amount(actual[0]) != update["values"][0][0]:
                raise ValueError(f"写入后回读不一致：{update['range']}")
        client.write_values(spreadsheet_token, [total_update])
        total_read = client.read_range(spreadsheet_token, total_update["range"])
        last_shop_row = int(re.search(r"![A-Z]+(\d+)", total_update["range"])[1]) - 1
        expected_total = sum((Decimal(str(as_amount(row[0]))) for row in written[:last_shop_row - 2]
                              if row and row[0] is not None and str(row[0]).strip()), Decimal(0))
        if (not total_read or not total_read[0]
                or abs(Decimal(str(as_amount(total_read[0][0]))) - expected_total) > Decimal("0.005")):
            raise ValueError(f"汇总回读校验失败：{total_update['range']}，预期 {expected_total}")
    return {
        "businessDate": business_date.isoformat(),
        "sourceFile": str(Path(source_file).resolve()),
        "filteredFile": str(filtered_path.resolve()),
        "sheetId": actual_sheet_id,
        "targetColumn": target_column,
        "dateHeaderCell": f"{target_column}2",
        "verified": not dry_run,
        "rowCount": len(rows),
        "updateCount": len(updates) + 1,
        "totalCell": total_update["range"].split("!")[1].split(":")[0],
        "totalFormula": total_update["values"][0][0],
        "dryRun": dry_run,
    }


def load_config(path: str | Path) -> dict[str, Any]:
    config_path = Path(path)
    if not config_path.exists():
        return {}
    payload = json.loads(config_path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("config.json 顶层必须是对象")
    return payload


def parse_business_date(value: str | None) -> date:
    if value:
        parsed = as_date(value)
        if parsed:
            return parsed
        raise ValueError(f"日期格式错误：{value}")
    return date.today() - timedelta(days=1)


def main() -> int:
    parser = argparse.ArgumentParser(description="筛选昨日销售并按店铺号写入润工作电子表格")
    parser.add_argument("--file", required=True, help="良域导出的 XLSX 文件")
    parser.add_argument("--date", help="业务日期，默认昨天，格式 YYYY-MM-DD")
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--api-base-url")
    parser.add_argument("--spreadsheet-token")
    parser.add_argument("--sheet-id")
    parser.add_argument("--output")
    parser.add_argument("--max-rows", type=int, default=2000)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    config = load_config(args.config)
    result = sync_file(
        args.file,
        parse_business_date(args.date),
        api_base_url=args.api_base_url or config.get("apiBaseUrl") or DEFAULT_API_BASE_URL,
        spreadsheet_token=args.spreadsheet_token or config.get("spreadsheetToken") or DEFAULT_SPREADSHEET_TOKEN,
        sheet_id=args.sheet_id or config.get("sheetId") or "",
        output_file=args.output,
        max_rows=args.max_rows,
        dry_run=args.dry_run,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
