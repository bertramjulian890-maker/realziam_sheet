from __future__ import annotations

import argparse
import json
import shutil
import tempfile
from calendar import monthrange
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from daily_sales_sheet import REQUIRED_HEADERS, as_amount, as_date, header_date


DEFAULT_TARGET = r"D:\git\realziam_sheet\2026年汕汇全场业绩手工表.xlsx"
TOTAL_LABELS = {"合计", "总计", "小计", "求和"}


def month_bounds(month: str) -> tuple[date, date]:
    try:
        start = date.fromisoformat(f"{month}-01")
    except ValueError as exc:
        raise ValueError(f"月份格式错误：{month}，需要 YYYY-MM") from exc
    if len(month) != 7:
        raise ValueError(f"月份格式错误：{month}，需要 YYYY-MM")
    return start, date(start.year, start.month, monthrange(start.year, start.month)[1])


def read_month_export(path: str | Path, month: str) -> dict[date, dict[str, int | float]]:
    start, end = month_bounds(month)
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            raise ValueError("导出表为空")
        positions = {str(value or "").strip(): index for index, value in enumerate(header)}
        missing = [field for field in REQUIRED_HEADERS if field not in positions]
        if missing:
            raise ValueError(f"导出表缺少字段：{', '.join(missing)}")

        by_day: dict[date, dict[str, int | float]] = {}
        for row_number, row in enumerate(iterator, start=2):
            sale_date = as_date(_cell(row, positions["销售日期"]))
            if sale_date is None or sale_date < start or sale_date > end:
                continue
            shop_id = str(_cell(row, positions["店铺号"]) or "").strip()
            if not shop_id:
                raise ValueError(f"导出表第 {row_number} 行店铺号为空")
            day_rows = by_day.setdefault(sale_date, {})
            if shop_id in day_rows:
                raise ValueError(f"{sale_date} 存在重复店铺号：{shop_id}")
            day_rows[shop_id] = as_amount(_cell(row, positions["销售总金额"]))
    finally:
        workbook.close()

    expected_days = [start + timedelta(days=offset) for offset in range((end - start).days + 1)]
    empty_days = [day.isoformat() for day in expected_days if not by_day.get(day)]
    if empty_days:
        raise ValueError(f"导出文件缺少整日数据（{len(empty_days)} 天）：{', '.join(empty_days)}")
    return by_day


def _cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _locate_month_columns(sheet: Any, month: str) -> dict[date, int]:
    start, end = month_bounds(month)
    columns: dict[date, int] = {}
    for column in range(12, sheet.max_column + 1):
        parsed = header_date(sheet.cell(2, column).value, start.year)
        if parsed and parsed.month == start.month:
            # Header year is intentionally ignored: this fixed annual template
            # can contain legacy Excel serials whose stored year is different.
            target = date(start.year, parsed.month, parsed.day)
            if target < start or target > end:
                continue
            if target in columns:
                raise ValueError(f"目标工作簿第 2 行存在重复日期列：{target}")
            columns[target] = column
    expected = [start + timedelta(days=offset) for offset in range((end - start).days + 1)]
    missing = [day.isoformat() for day in expected if day not in columns]
    if missing:
        raise ValueError(f"目标工作簿缺少 {month} 日期列：{', '.join(missing)}")
    return columns


def refresh_local_month(
    export_file: str | Path,
    target_file: str | Path = DEFAULT_TARGET,
    *,
    month: str,
    dry_run: bool = False,
) -> dict[str, Any]:
    export_path = Path(export_file).expanduser().resolve()
    target_path = Path(target_file).expanduser().resolve()
    if not export_path.is_file():
        raise FileNotFoundError(f"找不到导出文件：{export_path}")
    if not target_path.is_file():
        raise FileNotFoundError(f"找不到目标工作簿：{target_path}")
    by_day = read_month_export(export_path, month)

    workbook = load_workbook(target_path)
    try:
        sheet = workbook.active
        if [str(sheet.cell(1, col).value or "").strip() for col in (1, 2)] != ["店铺号", "店铺名称"]:
            raise ValueError("目标工作簿 A1、B1 必须分别为“店铺号”“店铺名称”")
        date_columns = _locate_month_columns(sheet, month)

        shop_rows: dict[str, int] = {}
        for row_number in range(3, sheet.max_row + 1):
            shop_id = str(sheet.cell(row_number, 1).value or "").strip()
            if not shop_id or shop_id in TOTAL_LABELS:
                continue
            if shop_id in shop_rows:
                raise ValueError(f"目标工作簿店铺号重复：{shop_id}")
            shop_rows[shop_id] = row_number
        if not shop_rows:
            raise ValueError("目标工作簿没有店铺数据行")
        last_shop_row = max(shop_rows.values())

        export_ids = {shop_id for daily in by_day.values() for shop_id in daily}
        unknown = sorted(export_ids - shop_rows.keys())
        if unknown:
            raise ValueError(f"以下导出店铺号在目标工作簿找不到：{', '.join(unknown[:30])}")

        changes: list[tuple[int, int, int | float | None]] = []
        day_totals: dict[str, int | float] = {}
        for business_date, column in sorted(date_columns.items()):
            amounts = by_day[business_date]
            for shop_id, row_number in shop_rows.items():
                changes.append((row_number, column, amounts.get(shop_id)))
            total = sum((Decimal(str(amount)) for amount in amounts.values()), Decimal(0))
            if total == total.to_integral_value():
                total_value: int | float = int(total)
            else:
                total_value = float(total)
            changes.append((last_shop_row + 1, column, total_value))
            day_totals[business_date.isoformat()] = total_value

        result: dict[str, Any] = {
            "month": month,
            "exportFile": str(export_path),
            "targetFile": str(target_path),
            "sheet": sheet.title,
            "days": len(date_columns),
            "shops": len(shop_rows),
            "updatedCells": len(changes),
            "clearedShopCells": sum(1 for row, col, value in changes
                                     if row <= last_shop_row and value is None),
            "totalRow": last_shop_row + 1,
            "totals": day_totals,
            "dryRun": dry_run,
        }
        if dry_run:
            return result

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = target_path.with_name(f"{target_path.stem}.backup_{timestamp}{target_path.suffix}")
        shutil.copy2(target_path, backup_path)
        temp_path: Path | None = None
        try:
            for row_number, column, value in changes:
                sheet.cell(row_number, column).value = value
            with tempfile.NamedTemporaryFile(
                prefix=f".{target_path.stem}_", suffix=target_path.suffix,
                dir=target_path.parent, delete=False,
            ) as temp_file:
                temp_path = Path(temp_file.name)
            workbook.save(temp_path)
            temp_path.replace(target_path)
        except Exception:
            if temp_path and temp_path.exists():
                temp_path.unlink()
            raise
        finally:
            workbook.close()

        check_book = load_workbook(target_path, read_only=True, data_only=True)
        try:
            check_sheet = check_book[sheet.title]
            for row_number, column, expected in changes:
                actual = check_sheet.cell(row_number, column).value
                if expected is None:
                    if actual not in (None, ""):
                        raise ValueError(f"本地写入回读校验失败：{check_sheet.cell(row_number, column).coordinate}")
                elif actual is None or Decimal(str(actual)) != Decimal(str(expected)):
                    raise ValueError(
                        f"本地写入回读校验失败：{check_sheet.cell(row_number, column).coordinate}，"
                        f"预期 {expected}，实际 {actual}"
                    )
        finally:
            check_book.close()

        result["backupFile"] = str(backup_path)
        result["verified"] = True
        return result
    finally:
        try:
            workbook.close()
        except Exception:
            pass


def main() -> int:
    parser = argparse.ArgumentParser(description="按指定月份把销售导出数据覆盖写入本地年度业绩工作簿")
    parser.add_argument("--month", required=True, help="业务月份 YYYY-MM，例如 2026-10")
    parser.add_argument("--export-file", required=True, help="良域导出的销售明细 XLSX")
    parser.add_argument("--target", default=DEFAULT_TARGET, help=f"本地目标工作簿，默认 {DEFAULT_TARGET}")
    parser.add_argument("--dry-run", action="store_true", help="只检查并预览，不写入目标工作簿")
    args = parser.parse_args()
    result = refresh_local_month(
        args.export_file, args.target, month=args.month, dry_run=args.dry_run
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
