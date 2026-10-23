"""Sales export through MMS HTTP APIs; no desktop or browser required."""
from __future__ import annotations

import argparse
import json
import logging
import os
import tempfile
import time
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import parse_qs, urlparse

from mms_ca_login import MMS, login_with_ca

LOGGER = logging.getLogger("daily_sales.export")
ROOT = Path(__file__).resolve().parent
EXPORT_PATH = "/api/mixc-inception-performance/v1/shops/sales-ledger/export-tasks"
TOKEN_PATH = "/api/bff-service/v1/fwshop/bps-common/v1/common/file/tokenUrl"
STORAGE_HOST = "storagecloud-api.saas.crland.com.cn"
DEFAULT_PROJECT = "20071"
DEFAULT_POSITION = "cd1f2432c46b4dbeac7b13ef4e1aaf87"


class _CaRequest:
    def __init__(self, session):
        self.session = session

    def _call(self, method, url, **kwargs):
        kwargs["timeout"] = kwargs.pop("timeout", 30000) / 1000
        if "data" in kwargs:
            kwargs["json"] = kwargs.pop("data")
        try:
            response = self.session.request(method, url, **kwargs)
        except Exception:
            raise RuntimeError("CA 认证网络请求失败，请检查内网及证书客户端") from None
        return SimpleNamespace(ok=response.ok, status=response.status_code,
                               url=response.url, json=response.json, text=lambda: response.text)

    def get(self, url, **kwargs):
        return self._call("GET", url, **kwargs)

    def post(self, url, **kwargs):
        return self._call("POST", url, **kwargs)


def project_settings(config):
    settings = config.get("mmsApi", {})
    if not isinstance(settings, dict):
        raise ValueError("mmsApi 必须为对象")
    source_project = parse_qs(urlparse(config.get("sourceUrl", "")).query).get("projectCode", [""])[0]
    project = str(settings.get("projectCode") or source_project or DEFAULT_PROJECT).strip()
    if source_project and project != source_project:
        raise ValueError("mmsApi.projectCode 与 sourceUrl 的项目编号不一致")
    position = str(settings.get("positionId") or os.environ.get("MMS_POSITION_ID", "")).strip()
    if not position and project == DEFAULT_PROJECT:
        position = DEFAULT_POSITION
    if not project or not position:
        raise ValueError("请配置 mmsApi.projectCode（或 sourceUrl 中的 projectCode）和 mmsApi.positionId（或 MMS_POSITION_ID）")
    return project, position


def create_session(project_code, position_id, auth=None):
    try:
        import requests
    except ImportError:
        raise RuntimeError("请安装 python -m pip install -r requirements-api.txt") from None
    session = requests.Session()
    session.headers.update({"Accept": "application/json", "Content-Type": "application/json"})
    try:
        auth = auth or {}
        name = (auth.get("caName") or os.environ.get("MMS_CA_NAME", "")).strip().upper()
        token = os.environ.get("MMS_ACCESS_TOKEN", "").strip()
        if not name:
            raise ValueError("请在 config.json 的 mmsApi.caName 中填写认证账号")
        if not token:
            LOGGER.info("开始 CA 接口认证")
            response = login_with_ca(SimpleNamespace(request=_CaRequest(session)), name=name,
                                     pin=auth.get("caPin"), identity_index=auth.get("certificateIndex"))
            data = response.get("data")
            token = data.get("token") if isinstance(data, dict) else None
        if not isinstance(token, str) or not token.strip():
            raise RuntimeError("良域登录未返回有效 data.token")
        session.headers.update({"Authorization": "Bearer " + token.removeprefix("Bearer "),
                                "loginName": name, "mixc-position-id": position_id,
                                "mixc-project-code": project_code, "projectCode": project_code})
        LOGGER.info("接口身份就绪；项目=%s", project_code)
        return session
    except Exception:
        session.close()
        raise


def _api(session, method, path, stage, timeout=30, **kwargs):
    try:
        response = session.request(method, MMS + path, timeout=timeout, allow_redirects=False, **kwargs)
        if response.status_code != 200:
            raise RuntimeError(f"{stage} HTTP {response.status_code}")
        body = response.json()
        if not isinstance(body, dict) or str(body.get("code")) != "0":
            raise RuntimeError(f"{stage} 业务响应失败；请检查登录和项目权限")
        return body.get("data")
    except RuntimeError:
        raise
    except Exception:
        # Signed URLs and Authorization must not enter exception logs.
        raise RuntimeError(f"{stage} 请求失败或响应不是有效 JSON") from None


def records(session, project, timeout=30):
    found = []
    page = 1
    while True:
        data = _api(session, "GET", EXPORT_PATH + "/records", "查询导出记录", timeout=timeout,
                    params={"pageIndex": page, "pageSize": 50, "projectNo": project,
                            "taskType": "PROJECT_SALES_EXPORT"})
        if not isinstance(data, dict) or not isinstance(data.get("records"), list):
            raise RuntimeError("导出记录响应缺少 records")
        rows = data["records"]
        if any(not isinstance(row, dict) or not row.get("businessId") for row in rows):
            raise RuntimeError("导出记录缺少 businessId，停止以避免下载旧文件")
        found.extend(rows)
        pages = int(data.get("pages", 1))
        if page >= pages:
            return found
        if not rows or page >= 100:
            raise RuntimeError("导出记录分页异常，停止取数")
        page += 1


def _storage_url(url):
    if not isinstance(url, str):
        raise RuntimeError("下载地址格式错误")
    parsed = urlparse(url)
    if (parsed.scheme != "https" or parsed.hostname != STORAGE_HOST
            or parsed.username or parsed.password or parsed.port not in (None, 443)):
        raise RuntimeError("下载地址不是已确认的存储云 HTTPS 地址")
    return url


def download_file(url, directory, start, end):
    import requests
    from openpyxl import load_workbook
    from daily_sales_sheet import REQUIRED_HEADERS

    url = _storage_url(url)
    directory = Path(directory).expanduser().resolve()
    directory.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        # Separate request prevents MMS bearer/position headers leaking to storage.
        with requests.get(url, stream=True, timeout=(15, 60), allow_redirects=False) as response:
            if response.status_code != 200:
                raise RuntimeError(f"文件下载 HTTP {response.status_code}；若跳转需核对下载链路")
            with tempfile.NamedTemporaryFile(dir=directory, suffix=".part.xlsx", delete=False) as output:
                temporary = Path(output.name)
                for chunk in response.iter_content(chunk_size=1024 * 1024):
                    if chunk:
                        output.write(chunk)
        workbook = load_workbook(temporary, read_only=True, data_only=True)
        try:
            headers = next(workbook.active.iter_rows(values_only=True), ())
            if not set(REQUIRED_HEADERS).issubset({str(value or "").strip() for value in headers}):
                raise RuntimeError("下载文件缺少销售必需字段")
        finally:
            workbook.close()
        filename = f"零售额台账_{start}_{end}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
        target = directory / filename
        temporary.replace(target)
        LOGGER.info("Excel 下载校验通过：%s", target)
        return target
    except RuntimeError:
        raise
    except Exception:
        raise RuntimeError("下载失败或文件不是有效销售 XLSX；已停止") from None
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def export_sales(*, project_code, position_id, start_date, end_date, download_dir,
                 timeout_seconds=180, session=None, poll_seconds=3, auth=None):
    start = date.fromisoformat(str(start_date)).isoformat()
    end = date.fromisoformat(str(end_date)).isoformat()
    if end < start or timeout_seconds <= 0:
        raise ValueError("日期范围或超时参数无效")
    owned = session is None
    session = session or create_session(project_code, position_id, auth=auth)
    try:
        before = {row["businessId"] for row in records(session, project_code)}
        payload = {"projectNo": project_code, "queryType": "DATE_TOTAL", "busiType": "",
                   "endDate": "", "shopFloor": "", "startDate": "", "workDateTime": "",
                   "buildingCode": "", "dealDateFrom": start, "dealDateTo": end,
                   "dataCaliber": "all", "taskType": "PROJECT_SALES_EXPORT",
                   "shopType": [], "userLoginNames": []}
        if _api(session, "POST", EXPORT_PATH + "/generator", "创建导出任务", json=payload) is not True:
            raise RuntimeError("创建导出任务未返回 data=true")
        LOGGER.info("导出已提交；项目=%s；日期=%s 至 %s", project_code, start, end)
        deadline = time.monotonic() + timeout_seconds
        selected = None
        while time.monotonic() < deadline:
            current = records(session, project_code, timeout=min(30, max(1, deadline - time.monotonic())))
            candidates = [row for row in current if row["businessId"] not in before
                          and row.get("projectNo") == project_code
                          and row.get("taskType") == "PROJECT_SALES_EXPORT"
                          and row.get("subTaskType") == "DATE_TOTAL"
                          and row.get("taskRangeBeginTime") == start
                          and row.get("taskRangeEndTime") == end]
            if len(candidates) > 1:
                raise RuntimeError("发现多条本次日期的新任务，无法确认归属，停止下载")
            if candidates:
                row = candidates[0]
                if selected and selected != row["businessId"]:
                    raise RuntimeError("本次导出任务标识变化，停止下载")
                selected = row["businessId"]
                state = row.get("taskState")
                LOGGER.info("本次导出状态=%s", state)
                if state in {"FAIL", "FAILED", "ERROR", "CANCELLED", "CANCELED"}:
                    raise RuntimeError("服务器导出任务失败或已取消")
                if state == "FINISH":
                    unsigned = _storage_url(row.get("downloadUrl"))
                    signed = _api(session, "GET", TOKEN_PATH, "获取下载凭证", params={"url": unsigned})
                    return download_file(_storage_url(signed), download_dir, start, end)
            time.sleep(min(poll_seconds, max(0, deadline - time.monotonic())))
        raise TimeoutError("等待本次导出完成超时；未下载历史文件")
    finally:
        if owned:
            session.close()


def main():
    from daily_sales_sheet import load_config
    from run_daily_sales import configure_logging, resolve_business_dates
    parser = argparse.ArgumentParser(description="通过良域接口取数，只下载 Excel，不写云表")
    parser.add_argument("--config", default=str(ROOT / "config.json"))
    parser.add_argument("--month")
    parser.add_argument("--date", help="指定业务日期 YYYY-MM-DD，只取当天")
    parser.add_argument("--month-to-yesterday", action="store_true")
    parser.add_argument("--start-date")
    parser.add_argument("--end-date")
    parser.add_argument("--project-code")
    parser.add_argument("--position-id")
    parser.add_argument("--timeout", type=int, default=180)
    args = parser.parse_args()
    configure_logging(ROOT)
    try:
        config = load_config(args.config)
        if args.project_code:
            config = dict(config, sourceUrl="", mmsApi={**config.get("mmsApi", {}), "projectCode": args.project_code})
        if args.position_id:
            config["mmsApi"] = {**config.get("mmsApi", {}), "positionId": args.position_id}
        project, position = project_settings(config)
        dates = resolve_business_dates(args.date, args.month_to_yesterday, args.start_date, args.end_date, month=args.month)
        path = export_sales(project_code=project, position_id=position, start_date=dates[0], end_date=dates[-1],
                            download_dir=config.get("downloadDir") or Path.home() / "Downloads",
                            timeout_seconds=args.timeout, auth=config.get("mmsApi", {}))
        print(json.dumps({"sourceFile": str(path), "projectCode": project,
                          "startDate": str(dates[0]), "endDate": str(dates[-1]), "cloudWritten": False}, ensure_ascii=False))
        return 0
    except Exception:
        LOGGER.exception("接口取数失败")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
