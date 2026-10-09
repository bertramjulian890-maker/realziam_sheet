from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from local_month_refresh import refresh_local_month


def make_export(path: Path, *, omit_day: int | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["店铺号", "店铺名称", "销售日期", "销售总金额"])
    for day in range(1, 32):
        if day == omit_day:
            continue
        sale_date = date(2026, 10, day)
        sheet.append(["SHOP1", "门店一", sale_date, day * 10])
        if day % 2 == 0:
            sheet.append(["SHOP2", "门店二", sale_date, 5])
    workbook.save(path)


def make_target(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "店铺号"
    sheet["B1"] = "店铺名称"
    sheet["C1"] = "保留字段"
    sheet["C3"] = "不应修改"
    sheet["A3"] = "SHOP1"
    sheet["B3"] = "门店一"
    sheet["A4"] = "SHOP2"
    sheet["B4"] = "门店二"
    sheet["A5"] = "合计"
    for day in range(1, 32):
        column = day + 11
        sheet.cell(1, column).value = f"星期{day}"
        # Template serial dates may have a legacy year; month/day is the key.
        sheet.cell(2, column).value = date(2025, 10, day)
        sheet.cell(3, column).value = -1
        sheet.cell(4, column).value = -2
        sheet.cell(5, column).value = -3
    workbook.save(path)


class LocalMonthRefreshTest(unittest.TestCase):
    def test_dry_run_previews_without_mutating_and_live_run_overwrites_month(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            export_file = root / "export.xlsx"
            target_file = root / "target.xlsx"
            make_export(export_file)
            make_target(target_file)

            today = date(2026, 10, 9)
            preview = refresh_local_month(
                export_file, target_file, month="2026-10", dry_run=True, today=today
            )
            self.assertEqual(preview["startDate"], "2026-10-01")
            self.assertEqual(preview["endDate"], "2026-10-08")
            self.assertEqual(preview["days"], 8)
            self.assertEqual(preview["shops"], 2)
            self.assertFalse(preview.get("verified", False))
            before = load_workbook(target_file, data_only=True)
            self.assertEqual(before.active["L3"].value, -1)
            before.close()

            result = refresh_local_month(export_file, target_file, month="2026-10", today=today)
            self.assertTrue(Path(result["backupFile"]).is_file())
            saved = load_workbook(target_file, data_only=True)
            sheet = saved.active
            self.assertEqual(sheet["L3"].value, 10)
            self.assertIsNone(sheet["L4"].value)  # Absent shop is cleared.
            self.assertEqual(sheet["L5"].value, 10)
            self.assertEqual(sheet["M3"].value, 20)
            self.assertEqual(sheet["M4"].value, 5)
            self.assertEqual(sheet["M5"].value, 25)
            self.assertEqual(sheet["S3"].value, 80)
            self.assertEqual(sheet["S5"].value, 85)
            self.assertEqual(sheet["T3"].value, -1)  # Today/future columns stay untouched.
            self.assertEqual(sheet["T5"].value, -3)
            self.assertEqual(sheet["C3"].value, "不应修改")
            saved.close()

    def test_missing_day_fails_before_target_changes_or_backup(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            export_file = root / "export.xlsx"
            target_file = root / "target.xlsx"
            make_export(export_file, omit_day=7)
            make_target(target_file)

            with self.assertRaisesRegex(ValueError, "缺少整日数据.*2026-10-07"):
                refresh_local_month(export_file, target_file, month="2026-10", today=date(2026, 10, 9))

            unchanged = load_workbook(target_file, data_only=True)
            self.assertEqual(unchanged.active["L3"].value, -1)
            unchanged.close()
            self.assertEqual(list(root.glob("target.backup_*.xlsx")), [])

    def test_unmatched_shop_fails_without_overwriting_target(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            export_file = root / "export.xlsx"
            target_file = root / "target.xlsx"
            make_export(export_file)
            source = load_workbook(export_file)
            source.active.append(["UNKNOWN", "未知门店", date(2026, 10, 1), 100])
            source.save(export_file)
            source.close()
            make_target(target_file)

            with self.assertRaisesRegex(ValueError, "目标工作簿找不到.*UNKNOWN"):
                refresh_local_month(export_file, target_file, month="2026-10", today=date(2026, 10, 9))

            unchanged = load_workbook(target_file, data_only=True)
            self.assertEqual(unchanged.active["L3"].value, -1)
            unchanged.close()

    def test_current_month_stops_yesterday_but_past_month_is_full(self) -> None:
        from local_month_refresh import month_bounds

        self.assertEqual(
            month_bounds("2026-10", today=date(2026, 10, 9)),
            (date(2026, 10, 1), date(2026, 10, 8)),
        )
        self.assertEqual(
            month_bounds("2026-09", today=date(2026, 10, 9)),
            (date(2026, 9, 1), date(2026, 9, 30)),
        )


if __name__ == "__main__":
    unittest.main()
