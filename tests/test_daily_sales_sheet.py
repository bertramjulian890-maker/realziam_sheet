from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


class DailySalesSheetTest(unittest.TestCase):
    def test_sync_uses_fixed_rows_and_checks_written_amount(self) -> None:
        from daily_sales_sheet import sync_file
        item = {"店铺号": "DL4101", "销售总金额": 45}
        with tempfile.TemporaryDirectory() as temp, \
                patch("daily_sales_sheet.filter_export", return_value=[item]), \
                patch("daily_sales_sheet.save_filtered_workbook", return_value=Path(temp) / "out.xlsx"), \
                patch("daily_sales_sheet.RworkSheetsClient") as client_type:
            client = client_type.return_value
            client.read_range.side_effect = [
                [["店铺号", "店铺名称"], [None] * 11 + ["9月14日"]],
                [["DL4101"], ["OTHER"], ["合计"]],
                [[45], [10.5], ["=SUM(L3:L4)"]],
                [[55.5]],
            ]
            result = sync_file("input.xlsx", date(2026, 9, 14),
                               api_base_url="http://test", spreadsheet_token="token", sheet_id="s")
            self.assertEqual(result["targetColumn"], "L")
            self.assertTrue(result["verified"])
            self.assertEqual(result["totalAmount"], 55.5)
            self.assertEqual(client.write_values.call_args_list[0].args, ("token", [
                {"range": "s!L3:L3", "values": [[45]]},
            ]))
            client.write_values.assert_called_with("token", [
                {"range": "s!L5:L5", "values": [[55.5]]},
            ])
            self.assertEqual(client.read_range.call_args_list[1].args[1], "s!A3:A2000")

    def test_total_is_below_last_cloud_shop_even_if_not_in_export(self) -> None:
        from daily_sales_sheet import plan_total
        ids = [f"shop-{i}" for i in range(99)] + ["合计", None]
        self.assertEqual(plan_total("s", "Z", ids),
                         {"range": "s!Z102:Z102", "values": [["=SUM(Z3:Z101)"]]})

    def test_filters_yesterday_and_writes_only_four_columns(self) -> None:
        from daily_sales_sheet import filter_export, save_filtered_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "export.xlsx"
            output = Path(temp_dir) / "filtered.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["其他", "店铺号", "店铺名称", "销售日期", "销售总金额"])
            sheet.append(["x", "EIGB101N01", "店铺甲", date(2026, 9, 14), 120.5])
            sheet.append(["y", "DL4101", "店铺乙", "2026/9/15", 400])
            sheet.append(["z", "CA11-005Z001", "店铺丙", "2026-09-14", 0])
            workbook.save(source)

            rows = filter_export(source, date(2026, 9, 14))
            save_filtered_workbook(rows, output)
            result = load_workbook(output, read_only=True, data_only=True)
            values = list(result.active.iter_rows(values_only=True))
            result.close()

        self.assertEqual(rows, [
            {"店铺号": "EIGB101N01", "店铺名称": "店铺甲", "销售日期": "2026-09-14", "销售总金额": 120.5},
            {"店铺号": "CA11-005Z001", "店铺名称": "店铺丙", "销售日期": "2026-09-14", "销售总金额": 0},
        ])
        self.assertEqual(values[0], ("店铺号", "店铺名称", "销售日期", "销售总金额"))
        self.assertEqual(values[1], ("EIGB101N01", "店铺甲", "2026-09-14", 120.5))

    def test_duplicate_shop_for_same_day_is_rejected(self) -> None:
        from daily_sales_sheet import filter_export

        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "export.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["店铺号", "店铺名称", "销售日期", "销售总金额"])
            sheet.append(["DL4101", "甲", "2026/9/14", 100])
            sheet.append(["DL4101", "甲", "2026/9/14", 200])
            workbook.save(source)
            with self.assertRaisesRegex(ValueError, "重复店铺号.*DL4101"):
                filter_export(source, date(2026, 9, 14))

    def test_matches_shop_id_and_existing_date_column(self) -> None:
        from daily_sales_sheet import plan_updates

        headers = [["店铺号", "店铺名称"], [None] * 11 + [f"9月{day}日" for day in range(1, 20)]]
        updates = plan_updates(
            "sheet-1",
            headers,
            ["CA11-005Z001", "DL4101", "EIGB101N01", "合计"],
            [{"店铺号": "EIGB101N01", "店铺名称": "甲", "销售日期": "2026-09-14", "销售总金额": 120.5}],
            date(2026, 9, 14),
        )
        self.assertEqual(updates, [{"range": "sheet-1!Y5:Y5", "values": [[120.5]]}])

    def test_missing_date_stops_without_creating_header(self) -> None:
        from daily_sales_sheet import plan_updates

        with self.assertRaisesRegex(ValueError, "固定模板不新增列"):
            plan_updates(
            "sheet-1",
            [["店铺号", "店铺名称"], [None] * 11 + ["9月14日"]],
            ["DL4101"],
            [{"店铺号": "DL4101", "店铺名称": "甲", "销售日期": "2026-09-15", "销售总金额": 45}],
            date(2026, 9, 15),
        )

    def test_cloud_date_serial_and_explicit_year(self) -> None:
        from daily_sales_sheet import header_date
        target = date(2026, 9, 14)
        serial = (target - date(1899, 12, 30)).days
        for value in (serial, float(serial), "9月14日", "2026/9/14"):
            self.assertEqual(header_date(value, 2026), target)
        self.assertNotEqual(header_date("2025/9/14", 2026), target)

    def test_actual_2025_serial_headers_match_2026_by_month_day(self) -> None:
        from daily_sales_sheet import plan_updates
        headers = [["店铺号", "店铺名称"], [None] * 11 + list(range(45901, 45931))]
        rows = [{"店铺号": "DL4101", "销售总金额": 45}]
        self.assertEqual(plan_updates("s", headers, ["DL4101", "合计"], rows, date(2026, 9, 15)),
                         [{"range": "s!Z3:Z3", "values": [[45]]}])
        headers[1].append("2026/9/15")
        with self.assertRaisesRegex(ValueError, "多个"):
            plan_updates("s", headers, ["DL4101"], rows, date(2026, 9, 15))


if __name__ == "__main__":
    unittest.main()
